#!/usr/bin/env python3
"""
Regenera src/datos/seed/*.ts desde el Excel.

El Excel es la especificación: ningún número del seed se escribe a mano.
Si cambias el Excel, corre esto y vuelve a probar:

    pip install openpyxl
    python3 herramientas/generar_seed.py
    npm test

Después de regenerar, sube VERSION_SEED en src/datos/db.ts para que las
instalaciones existentes vuelvan a sembrar los catálogos.
"""
import io
import json
import re
import sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
EXCEL = RAIZ / "Sistema_Mena_Plan_Integral.xlsx"
SEED = RAIZ / "src" / "datos" / "seed"

# Los 12 alimentos que más se usan salen arriba en el buscador.
FAVORITOS = {
    "Pechuga de pollo", "Clara de huevo", "Huevo entero", "Arroz blanco",
    "Avena en hojuelas", "Proteína whey", "Tilapia", "Batata", "Aguacate",
    "Aceite de oliva", "Banana (guineo)", "Vegetales salteados mixtos",
}

# Ejercicio (hoja GYM) -> columna de la hoja PROGRESION.
CLAVE_PROGRESION = {
    "press de banca con barra": "pressBancaTope",
    "remo en maquina con pecho apoyado": "remoMaquina",
    "prensa de piernas": "prensaPiernas",
    "empuje de cadera con barra (hip thrust)": "empujeCadera",
    "peso muerto rumano con barra hexagonal o mancuernas": "rdlHex",
    "dominadas lastradas": "dominadas",
    "pechadas hasta 2 reps del fallo": "pechadas",
}

MAPA_COMIDA = {
    "Desayuno": "desayuno", "Almuerzo": "almuerzo", "Merienda": "merienda",
    "Cena": "cena", "Cena ruptura": "cena", "Durante el día": "desayuno",
}

NOTAS_MENU = {
    ("ENTRENO", "desayuno"): "3 huevos + 6 claras. Avena pesada SECA.",
    ("ENTRENO", "almuerzo"): "Pollo CRUDO, arroz SECO.",
    ("ENTRENO", "merienda"): "1 scoop + 2 rebanadas.",
    ("ENTRENO", "cena"): "Tilapia ⇄ lubina ⇄ dorada ⇄ pollo, mismo peso.",
    ("AYUNO", "cena"): "Cena de ruptura ≈24 h después de la cena del jueves. Comida normal, no atracón.",
}

DIAS = {"LUNES": 1, "MARTES": 2, "MIÉRCOLES": 3, "JUEVES": 4, "VIERNES": 5, "SÁBADO": 6, "DOMINGO": 7}


def js(v):
    return json.dumps(v, ensure_ascii=False)


def opt(v):
    return "null" if v is None else (js(v) if isinstance(v, str) else str(v))


def normalizar(s):
    s = re.sub(r"^[abc][12]\s+", "", s.lower().strip())
    for a, b in zip("áéíóú", "aeiou"):
        s = s.replace(a, b)
    return s


def entero(v):
    if isinstance(v, int):
        return v
    if isinstance(v, str) and v.strip().isdigit():
        return int(v.strip())
    return None


def numero(v):
    try:
        n = float(v)
    except (TypeError, ValueError):
        return None
    return int(n) if n == int(n) else n


def segundos(txt):
    if not txt:
        return None
    t = str(txt).strip().lower()
    m = re.match(r"([\d.]+)\s*min", t)
    if m:
        return int(float(m.group(1)) * 60)
    m = re.match(r"([\d.]+)\s*s", t)
    return int(float(m.group(1))) if m else None


def escribir(nombre, contenido):
    (SEED / nombre).write_text(contenido, encoding="utf-8")
    print(f"  {nombre}")


def alimentos(wb):
    filas, i = [], 0
    for r in wb["ALIMENTOS"].iter_rows(min_row=4, values_only=True):
        if r[0] is None:
            continue
        i += 1
        filas.append(dict(
            id=i, nombre=r[0], estado=r[1] or "—", kcal=r[2], prot=r[3], carb=r[4],
            grasa=r[5], medida=r[6] or "", gramos=r[7] or 0, nota=r[8] or "",
            favorito=r[0] in FAVORITOS,
        ))
    out = io.StringIO()
    out.write("import type { Alimento } from '../../dominio/tipos'\n\n")
    out.write(f"/** Hoja ALIMENTOS del Excel — {len(filas)} filas, valores por 100 g. */\n")
    out.write("export const ALIMENTOS: Alimento[] = [\n")
    for a in filas:
        out.write(
            f"  {{ id: {a['id']}, nombre: {js(a['nombre'])}, estado: {js(a['estado'])}, "
            f"kcal100g: {a['kcal']}, prot100g: {a['prot']}, carb100g: {a['carb']}, "
            f"grasa100g: {a['grasa']}, medidaComun: {js(a['medida'])}, "
            f"gramosPorMedida: {a['gramos']}, nota: {js(a['nota'])}, "
            f"favorito: {'true' if a['favorito'] else 'false'} }},\n"
        )
    out.write("]\n")
    escribir("alimentos.ts", out.getvalue())
    return {a["nombre"]: a["id"] for a in filas}


def ejercicios(wb):
    lista, dia, orden, eid = [], None, 0, 0
    for r in wb["GYM"].iter_rows(values_only=True):
        if isinstance(r[0], str) and r[0].startswith("DÍA "):
            dia, orden = int(r[0].split()[1]), 0
            continue
        if dia is None or not isinstance(r[0], int):
            continue
        nombre = (r[1] or "").strip()
        if not nombre or normalizar(nombre).startswith(("protocolo articular", "sesion corta")):
            continue
        nota = (r[7] or "").strip()
        condicional = "CONDICIONAL" in nota
        orden += 1
        eid += 1
        lista.append(dict(
            id=eid, nombre=nombre, dia=dia, orden=orden,
            series=entero(r[2]), reps=str(r[3]).strip() if r[3] is not None else "",
            rpe=numero(r[4]), descanso=segundos(r[5]),
            # Si la nota ES la restricción, no se repite abajo en la pantalla.
            nota="" if condicional else nota,
            ancla="LEVANTAMIENTO ANCLA" in nota,
            restriccion=nota if condicional else None,
            clave=CLAVE_PROGRESION.get(normalizar(nombre)),
        ))
    out = io.StringIO()
    out.write("import type { Ejercicio } from '../../dominio/tipos'\n\n")
    out.write("/**\n * Hoja GYM — el split de 4 días, tal cual.\n")
    out.write(' * `esAncla` sale de la propia hoja: la nota técnica dice "TU LEVANTAMIENTO ANCLA".\n')
    out.write(" * `claveProgresion` enlaza con la hoja PROGRESION para precargar el peso de la semana.\n */\n")
    out.write("export const EJERCICIOS: Ejercicio[] = [\n")
    for e in lista:
        out.write(f"  {{\n    id: {e['id']}, nombre: {js(e['nombre'])}, dia: {e['dia']}, orden: {e['orden']},\n")
        out.write(f"    seriesPlan: {opt(e['series'])}, repsPlan: {js(e['reps'])}, "
                  f"rpePlan: {opt(e['rpe'])}, descansoSeg: {opt(e['descanso'])},\n")
        out.write(f"    notaTecnica: {js(e['nota'])},\n")
        out.write(f"    esAncla: {'true' if e['ancla'] else 'false'}, "
                  f"restriccion: {opt(e['restriccion'])}, claveProgresion: {opt(e['clave'])},\n  }},\n")
    out.write("]\n")
    escribir("ejercicios.ts", out.getvalue())


def progresion(wb):
    ws = wb["PROGRESION"]
    filas = [
        dict(semana=r[0], bloque=int(r[1]), cols=[str(c) for c in r[2:11]])
        for r in ws.iter_rows(values_only=True)
        if isinstance(r[0], int) and r[1] is not None and str(r[1]).isdigit()
    ]
    campos = ["pressBancaTope", "bancaBackoff", "remoMaquina", "prensaPiernas",
              "empujeCadera", "rdlHex", "dominadas", "pechadas", "enfoque"]
    out = io.StringIO()
    out.write("import type { CargaSemana } from '../../dominio/tipos'\n\n")
    out.write("/** Hoja PROGRESION — 12 semanas de cargas objetivo, calculadas desde los PR reales. */\n")
    out.write("export const PROGRESION: CargaSemana[] = [\n")
    for f in filas:
        pares = ", ".join(f"{k}: {js(v)}" for k, v in zip(campos, f["cols"]))
        out.write(f"  {{ semana: {f['semana']}, bloque: {f['bloque']}, {pares} }},\n")
    out.write("]\n\n")

    prs, capturando = [], False
    for r in ws.iter_rows(values_only=True):
        if r[0] and "TUS PR DECLARADOS" in str(r[0]):
            capturando = True
            continue
        if capturando:
            if r[0] and "CALENDARIO" in str(r[0]):
                break
            if r[1] and r[1] != "Levantamiento" and r[2]:
                prs.append((str(r[1]), str(r[2]), str(r[3]), str(r[10] or "")))
    out.write("export interface PrDeclarado { levantamiento: string; pr: string; rm1Estimado: string; nota: string }\n\n")
    out.write("/** PR declarados el 03/09/2026 — la base de toda la progresión. */\n")
    out.write("export const PR_DECLARADOS: PrDeclarado[] = [\n")
    for l, pr, rm, nota in prs:
        out.write(f"  {{ levantamiento: {js(l)}, pr: {js(pr)}, rm1Estimado: {js(rm)}, nota: {js(nota)} }},\n")
    out.write("]\n")
    escribir("progresion.ts", out.getvalue())


def menus(wb, ids):
    tipo_dia, agrupados = None, {}
    for r in wb["MENUS"].iter_rows(values_only=True):
        a = str(r[0]) if r[0] else ""
        if a.startswith("DÍA ENTRENO"):
            tipo_dia = "ENTRENO"
        elif a.startswith("DÍA LIGERO"):
            tipo_dia = "LIGERO"
        elif a.startswith("DÍA AYUNO"):
            tipo_dia = "AYUNO"
        elif tipo_dia and a in MAPA_COMIDA and r[1] in ids and isinstance(r[2], (int, float)) and r[2] > 0:
            clave = (tipo_dia, MAPA_COMIDA[a])
            agrupados.setdefault(clave, {"items": [], "etiqueta": a})
            agrupados[clave]["items"].append((ids[r[1]], r[2]))

    nombre_dia = {"ENTRENO": "día entreno", "LIGERO": "día ligero", "AYUNO": "viernes"}
    out = io.StringIO()
    out.write("import type { MenuGuardado } from '../../dominio/tipos'\n\n")
    out.write("/**\n * Hoja MENUS — los tres días modelo con sus gramos exactos.\n")
    out.write(" * Cargar uno de estos es lo que hace que registrar una comida tome 10 segundos.\n */\n")
    out.write("export const MENUS: MenuGuardado[] = [\n")
    for i, ((td, tipo), v) in enumerate(agrupados.items(), 1):
        items = ", ".join(f"{{ alimentoId: {a}, gramos: {g} }}" for a, g in v["items"])
        out.write(f"  {{\n    id: {i}, nombre: {js(v['etiqueta'] + ' ' + nombre_dia[td])}, "
                  f"tipoDia: {js(td)}, tipo: {js(tipo)},\n")
        out.write(f"    items: [{items}],\n    nota: {js(NOTAS_MENU.get((td, tipo), ''))},\n  }},\n")
    out.write("]\n")
    escribir("menus.ts", out.getvalue())


def horario(wb):
    bloques, dia, bid = [], None, 0
    for r in wb["HORARIO"].iter_rows(values_only=True):
        a = str(r[0]).strip() if r[0] else ""
        if a in DIAS:
            dia = DIAS[a]
            continue
        if not dia or len(a) != 5 or a[2] != ":":
            continue
        bid += 1
        bloques.append(dict(id=bid, dia=dia, hora=a, actividad=str(r[1] or ""),
                            detalle=str(r[2] or ""), bloque=str(r[3] or "")))
    out = io.StringIO()
    out.write("import type { BloqueHorario } from '../../dominio/tipos'\n\n")
    out.write("/** Hoja HORARIO — la semana bloque por bloque, cuadrada contra el calendario real. */\n")
    out.write("export const HORARIO: BloqueHorario[] = [\n")
    for b in bloques:
        out.write(f"  {{ id: {b['id']}, dia: {b['dia']}, hora: {js(b['hora'])}, "
                  f"actividad: {js(b['actividad'])}, detalle: {js(b['detalle'])}, "
                  f"bloque: {js(b['bloque'])} }},\n")
    out.write("]\n")
    escribir("horario.ts", out.getvalue())


def liturgico(wb):
    ws = wb["LITURGICO"]
    eventos, eid = [], 0
    for r in ws.iter_rows(min_row=3, values_only=True):
        if not r[0] or r[0] == "Frecuencia" or str(r[0]).startswith(("REGLA", "Si ")):
            continue
        if r[1] is None and r[2] is None:
            continue
        eid += 1
        nota = str(r[3] or "")
        eventos.append(dict(id=eid, nombre=nota.split(".")[0] if nota else str(r[0]),
                            frecuencia=str(r[0]), diaHora=str(r[1] or ""),
                            lugar=str(r[2] or ""), nota=nota))
    out = io.StringIO()
    out.write("import type { EventoLiturgico } from '../../dominio/tipos'\n\n")
    out.write("/** Hoja LITURGICO — compromisos fijos y rotativos. */\n")
    out.write("export const LITURGICO: EventoLiturgico[] = [\n")
    for e in eventos:
        out.write(f"  {{ id: {e['id']}, nombre: {js(e['nombre'])}, frecuencia: {js(e['frecuencia'])}, "
                  f"diaHora: {js(e['diaHora'])}, lugar: {js(e['lugar'])}, nota: {js(e['nota'])} }},\n")
    out.write("]\n\n")
    out.write("/** Regla de conflicto, textual: gana el compromiso espiritual. */\n")
    out.write("export const REGLA_CONFLICTO = [\n")
    for r in ws.iter_rows(values_only=True):
        if r[0] and str(r[0]).startswith("Si "):
            out.write(f"  {js(str(r[0]))},\n")
    out.write("]\n")
    escribir("liturgico.ts", out.getvalue())


def main():
    if not EXCEL.exists():
        sys.exit(f"No encuentro {EXCEL}")
    wb = openpyxl.load_workbook(EXCEL)
    print(f"Sembrando desde {EXCEL.name}:")
    ids = alimentos(wb)
    ejercicios(wb)
    progresion(wb)
    menus(wb, ids)
    horario(wb)
    liturgico(wb)
    print("Listo. Corre `npm test` y sube VERSION_SEED en src/datos/db.ts.")


if __name__ == "__main__":
    main()
